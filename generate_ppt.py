"""
Final comprehensive PPT editor for the NLP Course Project.
Replaces ALL Transcripto content with CourseClassifier content.
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation

TEMPLATE = r"C:\Users\vedan\.gemini\antigravity\brain\ecd83f14-df7c-4070-9494-bd179877ebe3\.user_uploaded\media_1791389416348.pptx"
OUTPUT = r"d:\nlp project\NLP_Course_Project_Presentation.pptx"


def deep_replace_all(shape, old, new):
    """Replace ALL occurrences of old with new in shape and all children, across all paragraphs."""
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            if old in para.text:
                new_text = para.text.replace(old, new)
                if para.runs:
                    para.runs[0].text = new_text
                    for r in para.runs[1:]:
                        r.text = ""
    if shape.shape_type == 6:
        for child in shape.shapes:
            deep_replace_all(child, old, new)


def slide_replace(slide, old, new):
    for shape in slide.shapes:
        deep_replace_all(shape, old, new)


def main():
    print("Loading template...")
    prs = Presentation(TEMPLATE)
    S = prs.slides

    # ═══════════════════════════════════════
    # GLOBAL replacements (apply to ALL slides)
    # ═══════════════════════════════════════
    global_replacements = [
        ("Transcripto", "CourseClassifier"),
        ("Meeting Minutes Generator", "Academic Discipline Predictor"),
        ("AI Meeting Minutes Generator", "Academic Course Description Classification"),
        ("meeting minutes", "course classification"),
        ("Meeting Transcript", "Course Description"),
        ("Structured Minutes", "Discipline Prediction"),
    ]

    for slide in S:
        for old, new in global_replacements:
            slide_replace(slide, old, new)

    # ═══════════════════════════════════════
    # SLIDE 1: Title
    # ═══════════════════════════════════════
    print("Slide 1: Title")
    s = S[0]
    slide_replace(s, "NLP-Based Automatic Meeting Minutes Generation",
                     "NLP-Based Academic Course Description Classification")
    slide_replace(s, "Shreeyansh Singh", "Vedant")
    slide_replace(s, "Roll 44", "Roll XX")
    slide_replace(s, "Parth Sitoot", "Team Member 2")
    slide_replace(s, "Roll 45", "Roll XX")
    slide_replace(s, "Keshavkumar Suthar", "Team Member 3")
    slide_replace(s, "Roll 49", "Roll XX")
    slide_replace(s, "Saiesh Sutar", "Team Member 4")
    slide_replace(s, "Roll 50", "Roll XX")
    slide_replace(s, "Guide - Prof. Sridhar Subramanian", "Guide - Prof. [Your Guide Name]")

    # ═══════════════════════════════════════
    # SLIDE 2: Introduction
    # ═══════════════════════════════════════
    print("Slide 2: Introduction")
    s = S[1]

    # Fix top-level text boxes
    slide_replace(s, "Written meeting", "University course")
    slide_replace(s, "transcripts", "descriptions")
    slide_replace(s, "MoM sections", "10 academic")
    slide_replace(s, "structured output", "disciplines classified")
    slide_replace(s, "model variants", "transformer models")
    slide_replace(s, "10K", "3K")
    slide_replace(s, "16K", "3K")
    slide_replace(s, "max input", "course descriptions")
    slide_replace(s, "tokens", "in dataset")
    slide_replace(s, "100", "3,000")
    slide_replace(s, "9", "10")

    # Fix the quote text
    slide_replace(s, "A meeting is only useful when the next person can act on it",
                     "Automating course classification enables smarter curriculum planning")

    # Fix nested descriptive paragraphs
    slide_replace(s, "The problem is not meeting capture.",
                     "Classifying courses into disciplines")
    slide_replace(s, "It is meeting memory.",
                     "is key to academic analytics.")
    slide_replace(s, "Long meetings hide decisions inside pages of dialogue.",
                     "University catalogs contain thousands of course descriptions.")
    slide_replace(s, "Owners, deadlines, numbers and caveats appear across speakers.",
                     "Manual classification is tedious and does not scale.")
    slide_replace(s, "Generic summaries can be fluent but lose operational structure.",
                     "We use 4 transformer models to auto-classify into 10 disciplines.")
    slide_replace(s, "turns a transcript into reviewable, structured",
                     "achieves 100% accuracy with millisecond")

    # ═══════════════════════════════════════
    # SLIDE 3: Problem Statement & Objectives
    # ═══════════════════════════════════════
    print("Slide 3: Problem Statement")
    s = S[2]
    slide_replace(s, "What we set out to solve", "What we set out to build")
    slide_replace(s,
        "Build a system that converts long written meeting transcripts into concise, structured minutes while preserving decisions, action items, participants, dates, numbers and source context.",
        "Build an automated NLP system that classifies university course descriptions into 10 academic disciplines using 4 transformer-based models from Hugging Face.")
    slide_replace(s, "Parse & Clean", "Data Collection")
    slide_replace(s,
        "Normalize transcript text, recover speaker turns and preserve conversation order.",
        "Curate 3,000 course descriptions from MIT OCW, Coursera, edX, Stanford, UC Berkeley & Harvard (300 per discipline).")
    slide_replace(s, "Focus Long Context", "Multi-Model Training")
    slide_replace(s,
        "Use TF-IDF, sentence salience and action/decision cues to retain important content.",
        "Fine-tune 4 Hugging Face transformers: BERT, RoBERTa, DeBERTa, and T5 for sequence classification.")
    slide_replace(s, "Generate Structured MoM", "Evaluate & Compare")
    slide_replace(s,
        "Fine-tuned LongT5 generates an 9-section", "Benchmark with Accuracy, F1, ROC-AUC across 10")
    slide_replace(s, "Validate & Evaluate", "Deploy & Test")
    slide_replace(s,
        "Compare models with ROUGE-1/2/L and BERTScore; add section, numeric and entity support diagnostics.",
        "Interactive Flask web app with model comparison dashboard. CLI testing with pre-configured test suite.")

    # ═══════════════════════════════════════
    # SLIDE 5: Research Gap
    # ═══════════════════════════════════════
    print("Slide 5: Research Gap")
    s = S[4]
    slide_replace(s, "Structured minutes vs. free-form summaries",
                     "Single-model vs. multi-architecture comparison")
    slide_replace(s,
        "Many systems optimize summary fluency, but operational fields can remain implicit.",
        "Most studies evaluate only one architecture; cross-architecture benchmarking on course text is absent.")
    slide_replace(s,
        "Use an 9-section MoM schema with decisions, actions, risks and next steps.",
        "Train and compare 4 transformer architectures (BERT, RoBERTa, DeBERTa, T5) on the same dataset.")
    slide_replace(s, "Long-context coverage & topic drift", "Diverse & balanced dataset")
    slide_replace(s,
        "Long meetings introduce omissions, repetition and multi-topic drift.",
        "Existing datasets are small, single-source, or class-imbalanced, limiting generalization.")
    slide_replace(s,
        "Use LongT5 plus salience-aware context selection for long transcripts.",
        "Build a 3,000-sample balanced dataset (300/class) from 6 university sources with 60/20/20 split.")
    slide_replace(s, "Evaluation beyond ROUGE", "Evaluation beyond accuracy")
    slide_replace(s,
        "Lexical overlap alone does not prove factual support or field completeness.",
        "Accuracy alone does not reveal per-class performance or calibration quality.")
    slide_replace(s,
        "Add section, numeric and entity support diagnostics to the benchmark.",
        "Add ROC-AUC curves (per-class + micro/macro), confusion matrices, and latency benchmarks.")

    # ═══════════════════════════════════════
    # SLIDE 7: Results header
    # ═══════════════════════════════════════
    print("Slide 7: Results")
    s = S[6]
    slide_replace(s, "LongT5 leads across all evaluation metrics",
                     "All 4 models achieve 100% accuracy across 10 disciplines")

    # ═══════════════════════════════════════
    # SLIDE 8: Base Paper comparison header
    # ═══════════════════════════════════════
    print("Slide 8: Base Paper Comparison")
    s = S[7]
    slide_replace(s, "Base Paper (LongT5) vs Our Work",
                     "Literature Survey Baselines vs Our Work")

    # ═══════════════════════════════════════
    # SLIDE 9: Implementation
    # ═══════════════════════════════════════
    print("Slide 9: Implementation")
    s = S[8]
    # Already handled by global replacements

    # ═══════════════════════════════════════
    # SLIDE 10: Conclusion
    # ═══════════════════════════════════════
    print("Slide 10: Conclusion")
    s = S[9]
    slide_replace(s,
        "Built an end-to-end 100-transcript pipeline from web cleaning and NLP analysis to structured",
        "Built an end-to-end pipeline classifying 3,000 course descriptions across 10 disciplines using 4 transformer")
    slide_replace(s,
        "Fine-tuned LongT5 delivered the strongest overall summary quality.",
        "All 4 fine-tuned transformers achieved 100% test accuracy. BERT selected as best model.")
    slide_replace(s, "ROUGE-1 F1 - 0.810", "Test Accuracy – 100.0%")
    slide_replace(s, "ROUGE-L F1 - 0.743", "Macro F1 – 1.0000")
    slide_replace(s, "BERTScore - 0.930", "ROC-AUC – 1.0000")

    # ═══════════════════════════════════════
    # SLIDE 12: References
    # ═══════════════════════════════════════
    print("Slide 12: References")
    s = S[11]

    # Replace each reference by matching unique author
    ref_replacements = [
        ("Kartik",
         "J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, «BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding,» Proc. NAACL-HLT 2019, pp. 4171–4186, 2019."),
        ("Chatterjee",
         "Y. Liu et al., «RoBERTa: A Robustly Optimized BERT Pretraining Approach,» arXiv preprint arXiv:1907.11692, 2019."),
        ("Agarwal",
         "P. He, X. Liu, J. Gao, and W. Chen, «DeBERTa: Decoding-enhanced BERT with Disentangled Attention,» Proc. ICLR 2021, 2021."),
        ("Xue",
         "C. Raffel et al., «Exploring the Limits of Transfer Learning with a Unified Text-to-Text Transformer,» JMLR, vol. 21, pp. 1–67, 2020."),
        ("Hu et al",
         "A. Vaswani et al., «Attention Is All You Need,» Advances in Neural Information Processing Systems (NeurIPS), 2017."),
        ("Dettmers",
         "T. Wolf et al., «Transformers: State-of-the-Art Natural Language Processing,» Proc. EMNLP 2020, pp. 38–45, 2020."),
    ]

    def deep_set_full_by_substr(shape, match_substr, new_text):
        if shape.has_text_frame:
            if match_substr in shape.text_frame.text:
                # Set entire text frame to new text (first paragraph, first run)
                for pi, para in enumerate(shape.text_frame.paragraphs):
                    if pi == 0:
                        if para.runs:
                            para.runs[0].text = new_text
                            for r in para.runs[1:]:
                                r.text = ""
                    else:
                        for r in para.runs:
                            r.text = ""
                return True
        if shape.shape_type == 6:
            for child in shape.shapes:
                if deep_set_full_by_substr(child, match_substr, new_text):
                    return True
        return False

    for match, new_text in ref_replacements:
        for shape in s.shapes:
            if deep_set_full_by_substr(shape, match, new_text):
                break

    # ═══════════════════════════════════════
    # Save
    # ═══════════════════════════════════════
    prs.save(OUTPUT)
    print(f"\n{'='*60}")
    print(f"  ✅ Presentation saved to: {OUTPUT}")
    print(f"{'='*60}")

    # ── Verify ──
    print("\n=== VERIFICATION ===")
    prs2 = Presentation(OUTPUT)

    def collect_all_text(shape):
        texts = []
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                if para.text.strip():
                    texts.append(para.text.strip()[:200])
        if shape.shape_type == 6:
            for child in shape.shapes:
                texts.extend(collect_all_text(child))
        return texts

    for i, slide in enumerate(prs2.slides):
        texts = []
        for shape in slide.shapes:
            texts.extend(collect_all_text(shape))
        if texts:
            print(f"\n--- Slide {i+1} ---")
            for t in texts:
                print(f"  {t}")


if __name__ == "__main__":
    main()
